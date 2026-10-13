"""Build templates/security-review-questionnaire.xlsx and .csv from docs/10-security-review-questionnaire.md."""
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

root = Path(__file__).resolve().parent.parent
src = (root / "docs" / "10-security-review-questionnaire.md").read_text()

rows, section = [], ""
for line in src.splitlines():
    m = re.match(r"## ([A-I])\. (.+)", line)
    if m:
        section = f"{m.group(1)}. {m.group(2)}"
        continue
    m = re.match(r"\| ([A-I]\d+) \| (.+?) \| (.+?) \|$", line)
    if m:
        rows.append([m.group(1), section, m.group(2), m.group(3)])

if not rows:
    sys.exit("no questions parsed")

head = ["ID", "Section", "Question", "Evidence to request", "Response", "Status", "Cloud / region", "Owner", "Source and date", "Notes"]
out = root / "templates"
out.mkdir(exist_ok=True)

with open(out / "security-review-questionnaire.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(head)
    for r in rows:
        w.writerow(r + [""] * (len(head) - len(r)))

wb = Workbook()
ws = wb.active
ws.title = "Questionnaire"
ws.append(head)
for r in rows:
    ws.append(r + [""] * (len(head) - len(r)))
fill = PatternFill("solid", fgColor="1B3139")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = fill
    c.alignment = Alignment(vertical="center", wrap_text=True)
for col, width in zip("ABCDEFGHIJ", [6, 22, 60, 34, 50, 16, 18, 16, 24, 40]):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)
ws.freeze_panes = "D2"
ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type="list", formula1='"Not started,In progress,Answered,Needs confirmation,Accepted risk,Blocked"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"F2:F{len(rows) + 1}")

guide = wb.create_sheet("How to use")
for line in [
    "Fill Response, Status, Cloud / region, Owner, and Source and date for each question.",
    "Cite a public document or a written statement from the vendor; record the date.",
    "Use 'Needs confirmation' for items that only a vendor can settle; see notes/open-questions.md.",
    "Questions come from docs/10-security-review-questionnaire.md. Regenerate with scripts/build_questionnaire_sheet.py.",
]:
    guide.append([line])
guide.column_dimensions["A"].width = 110
wb.save(out / "security-review-questionnaire.xlsx")
print(len(rows), "questions")
